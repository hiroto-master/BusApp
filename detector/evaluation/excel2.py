import os
import cv2
from ultralytics import YOLO
from datetime import datetime
import time
import psutil
import openpyxl

# --- 設定 ---
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
INPUT_FOLDER = os.path.join("yolo_9", 'input_images')
OUTPUT_FOLDER_PROCESSED = os.path.join("yolo_9", 'output_images2') # 検証用に出力先を分けました
EXCEL_FILE_PATH = os.path.join("yolo_9", 'people_count_log2.xlsx')

# 検証するモデル
YOLO_MODEL_PATH = 'yolov8n.pt'

TARGET_CLASSES = [0]  # 'person' (人物)

# --- フォルダの準備 ---
if not os.path.exists(OUTPUT_FOLDER_PROCESSED):
    os.makedirs(OUTPUT_FOLDER_PROCESSED)
    print(f"出力フォルダを作成しました: {OUTPUT_FOLDER_PROCESSED}")

# --- Excelへの書き込み関数 ---
def save_to_excel(timestamp, filename, count, conf_threshold):
    """
    検証結果をExcelに追記保存する関数
    A列: しきい値, B列: 日時, C列: ファイル名, D列: 人数
    """
    try:
        if os.path.exists(EXCEL_FILE_PATH):
            wb = openpyxl.load_workbook(EXCEL_FILE_PATH)
            ws = wb.active
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Experiment Log"
            # ヘッダー行
            ws.append(["Threshold", "Timestamp", "Filename", "Person Count", "Model"])
            print(f"新規Excelファイルを作成しました: {EXCEL_FILE_PATH}")

        # データを1行追記 (A列にしきい値を入れる)
        ws.append([conf_threshold, timestamp, filename, count, YOLO_MODEL_PATH])
        
        wb.save(EXCEL_FILE_PATH)
        # print(f"    Excelに記録しました") 

    except PermissionError:
        print("    【エラー】Excelファイルが開かれています。閉じてから再実行してください。")
        time.sleep(1) # エラー時は少し待つ
    except Exception as e:
        print(f"    Excel保存エラー: {e}")

def add_blank_row():
    """
    Excelに空行を1行追加する
    """
    try:
        if os.path.exists(EXCEL_FILE_PATH):
            wb = openpyxl.load_workbook(EXCEL_FILE_PATH)
            ws = wb.active
            ws.append([]) # 空のリストを追加すると空行になる
            wb.save(EXCEL_FILE_PATH)
    except Exception:
        pass

def run_experiment():
    """
    しきい値を1.0から0.1まで変えながら検証を実行する
    """
    print(f"YOLOモデルをロード中: {YOLO_MODEL_PATH}...")
    try:
        # ループの外でモデルを一度だけロード（効率化のため）
        model = YOLO(YOLO_MODEL_PATH)
        print("モデルのロードが完了しました。")
    except Exception as e:
        print(f"モデルのロードに失敗しました: {e}")
        return

    # 画像ファイルのリストを取得
    if not os.path.exists(INPUT_FOLDER):
        print(f"入力フォルダが見つかりません: {INPUT_FOLDER}")
        return

    image_files = [f for f in os.listdir(INPUT_FOLDER) if f.lower().endswith(('.png', '.jpg', '.jpeg'))]
    if not image_files:
        print("処理する画像が見つかりません。")
        return

    print(f"検証対象画像: {len(image_files)}枚")

    # --- ★しきい値を 10 (1.0) から 1 (0.1) までループ ---
    for i in range(10, 0, -1):
        # floatの計算誤差を防ぐため、整数で割る
        current_conf = i / 10.0
        
        print(f"\n{'='*40}")
        print(f"検証開始: しきい値 (Conf) = {current_conf}")
        print(f"{'='*40}")

        processed_count = 0

        for filename in sorted(image_files):
            filepath = os.path.join(INPUT_FOLDER, filename)
            # print(f"  処理中 ({current_conf}): {filename}")

            try:
                # 推論実行 (ここで動的にconfを指定)
                results = model(filepath, conf=current_conf, classes=TARGET_CLASSES, verbose=False)
                result = results[0]

                # 人数カウント
                person_boxes = [box for box in result.boxes if box.cls == 0]
                person_count = len(person_boxes)
                
                # Excelへ記録
                current_time_str = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                save_to_excel(current_time_str, filename, person_count, current_conf)

                # 画像保存 (ファイル名にしきい値を含める)
                # 例: processed_conf0.5_image01.jpg
                annotated_image = result.plot()
                output_filename = f"processed_conf{current_conf}_{filename}"
                output_filepath = os.path.join(OUTPUT_FOLDER_PROCESSED, output_filename)
                
                cv2.imwrite(output_filepath, annotated_image)
                
                processed_count += 1

            except Exception as e:
                print(f"  エラー発生: {filename} の処理中にエラー: {e}")
            add_blank_row()
        print(f"  >>> しきい値 {current_conf} の処理完了 ({processed_count}枚)")

    print("\nすべての検証が終了しました。")

if __name__ == "__main__":
    print("--- YOLO 検証モード ---")
    print(f"対象フォルダ: {INPUT_FOLDER}")
    print(str(YOLO_MODEL_PATH)+"で、しきい値を 1.0 -> 0.1 まで変化させて記録します。")
    print("開始するには Enter を押してください...")
    try:
        input()
        start_time = time.time()
        run_experiment()
        end_time = time.time()
        print(f"所要時間: {end_time - start_time:.1f}秒")
    except KeyboardInterrupt:
        print("\n中断しました。")
        exit()
